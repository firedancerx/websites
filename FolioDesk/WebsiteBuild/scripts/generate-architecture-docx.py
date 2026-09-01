import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    for run in h.runs:
        run.font.name = 'Calibri'
        if level == 1:
            run.font.size = Pt(17)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 118, 110) # Teal #0f766e
        elif level == 2:
            run.font.size = Pt(13.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 64, 175) # Blue #1e40af
        elif level == 3:
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85) # Slate #334155
    return h

def create_callout(doc, title, text, bg_hex="EFF6FF", border_hex="3B82F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.name = 'Calibri'
    run_title.font.bold = True
    run_title.font.size = Pt(11)
    run_title.font.color.rgb = RGBColor(30, 58, 138)
    
    run_text = p.add_run(text)
    run_text.font.name = 'Calibri'
    run_text.font.size = Pt(10.5)
    run_text.font.color.rgb = RGBColor(51, 65, 85)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def style_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if is_header:
                set_cell_shading(cell, "0F766E")
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    for r in p.runs:
                        r.font.name = 'Calibri'
                        r.font.bold = True
                        r.font.size = Pt(10)
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                set_cell_shading(cell, bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for r in p.runs:
                        r.font.name = 'Calibri'
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(30, 41, 59)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if col_widths and j < len(col_widths):
                cell.width = Inches(col_widths[j])

def build_docx():
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_eyebrow = title_p.add_run("FOLIODESK ENTERPRISE SYSTEM SPECIFICATION\n")
    run_eyebrow.font.name = 'Calibri'
    run_eyebrow.font.bold = True
    run_eyebrow.font.size = Pt(10)
    run_eyebrow.font.color.rgb = RGBColor(15, 118, 110)

    run_main_title = title_p.add_run("Affiliate Prospect Funnel, Review Cycles, Closure Period Policy & Payment Architecture")
    run_main_title.font.name = 'Calibri'
    run_main_title.font.bold = True
    run_main_title.font.size = Pt(21)
    run_main_title.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(4)
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("End-to-End Technical & Operational Architecture: Prospect Exclusivity, Step-by-Step Administrative Review Cycles, Global Closure Deadlines, Appeals, Snapshot Rate Locking, and Consolidated Batch Disbursements\nVersion 3.0 | Production Release | FolioDesk Enterprise Architecture")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("―" * 55)

    # SECTION 1: EXECUTIVE SUMMARY
    add_header_styled(doc, "1. Executive Summary & Core Architectural Principles", level=1)
    p = doc.add_paragraph(
        "The FolioDesk Affiliate Commercial Subsystem provides an integrated operating environment governing prospective client introductions, stage-by-stage sales funnel progression, administrative review cycles, closure period enforcement, management approval gating, and consolidated batch bank transfers. "
        "The architecture is engineered upon seven core pillars:"
    )
    p.paragraph_format.line_spacing = 1.15

    principles = [
        ("Anti-Poaching Exclusivity", "No affiliate can register or claim an existing prospect (by full registered company name) while it is actively logged by another affiliate until the active case is formally closed or stopped."),
        ("Interactive Step-by-Step Review Cycles", "Affiliates submit timestamped milestone updates along their sales funnel. Each update undergoes Administrative Review with the power to 'Acknowledge & Approve' or 'Return for Review (with Remarks)' for affiliate improvement and resubmission."),
        ("Configurable Global Closure Period", "Admin settings enforce a global negotiation window (e.g. 90 days) from the initial account log date. Exceeding this period empowers Admin to force close the attempt, triggering an affiliate appeal mechanism."),
        ("Extensible Negotiation Policies", "Admins can grant deadline extensions with or without an appeal, logging every time extension immutably in audit logs."),
        ("Snapshot Rate Locking at Submission", "When an invoice collection is recorded, active commission rates are captured as a snapshot and locked permanently to that collection, insulating historical payouts from future policy changes."),
        ("Mandatory Management Approval & Immutability", "Collections must be acknowledged and approved by Management before commission vouchers are issued. Approval instantly seals the transaction chain as strictly immutable (is_immutable = 1)."),
        ("Consolidated Batch Disbursements", "Pending commission vouchers are aggregated by Beneficiary Affiliate, allowing administrators to execute single 1-click bank transfers under a master voucher (DISB-2026-XXXXX).")
    ]
    for title, desc in principles:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(3)
        r_t = bp.add_run(f"{title}: ")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(15, 118, 110)
        r_d = bp.add_run(desc)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # SECTION 2: PROSPECT EXCLUSIVITY & ANTI-POACHING
    add_header_styled(doc, "2. Prospect Exclusivity & Anti-Poaching Registration Policy", level=1)
    doc.add_paragraph(
        "To protect affiliate commercial investment, lead generation efforts, and client relationships, the system enforces automated real-time anti-poaching exclusivity:"
    )
    create_callout(doc, "Deal Registration Exclusivity Rule",
                   "When an affiliate attempts to name and register a new prospect by company name:\n"
                   "1. The engine checks deal_pipeline for any active transaction matching LOWER(TRIM(customer_name)).\n"
                   "2. An active transaction is defined as any deal not in ABORTED or UNCOLLECTIBLE status, and where is_force_closed = 0.\n"
                   "3. If an active registration exists under another affiliate, the registration is rejected with the exact reason and existing affiliate assignment.\n"
                   "4. If a previous attempt was stopped or force-closed, the company name is released back into the public pool for new registrations.",
                   bg_hex="EFF6FF", border_hex="2563EB")

    # SECTION 3: SALES FUNNEL STEPS & ADMIN REVIEW CYCLES
    add_header_styled(doc, "3. Sales Funnel Step Logging & Administrative Review Cycles", level=1)
    doc.add_paragraph(
        "Rather than relying on opaque status flags, affiliates document the commercial evolution of their deals by submitting structured Funnel Steps into deal_funnel_steps:"
    )

    steps_desc = [
        ("Step Submission by Affiliate", "Affiliates select the target commercial stage (e.g. Qualified, Proposal Sent, Contract Signed), provide a step headline (e.g. 'Conducted Live Product Demo with Procurement Committee'), and record detailed progress notes."),
        ("Initial Gated State (PENDING_REVIEW)", "The step is logged with an immutable timestamp and placed in the administrative review queue."),
        ("Admin Action A: Acknowledge & Approve", "Admin verifies the milestone. The step status is updated to ACKNOWLEDGED, the deal's main status is automatically advanced to the target stage, and reviewer identity and timestamp are recorded."),
        ("Admin Action B: Return for Review (with Remarks)", "If information is insufficient (e.g. missing quotation copy or unclear tender timeline), Admin returns the step with mandatory corrective remarks. The step status becomes RETURNED_FOR_REVIEW."),
        ("Affiliate Resubmission", "The affiliate sees an alert banner on their prospect journey path, updates their notes, and resubmits for immediate re-evaluation.")
    ]
    for title, desc in steps_desc:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(3)
        r_t = bp.add_run(f"{title}: ")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(30, 64, 175)
        r_d = bp.add_run(desc)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # SECTION 4: CLOSURE PERIOD POLICY, APPEALS & EXTENSIONS
    add_header_styled(doc, "4. Global Closure Period Policy, Forced Closure & Appeals", level=1)
    doc.add_paragraph(
        "To prevent stale or abandoned deals from perpetually locking company names, the platform introduces a lifecycle closure management engine:"
    )

    closure_table_data = [
        ("Mechanism", "Operational Trigger & Rule", "Outcome / Impact"),
        ("Global Closure Period", "Configurable in /admin/settings (default 90 days from initial log date).", "Defines the baseline time allowed to close a customer contract."),
        ("Overdue Tracking", "Calculated as: created_at + closure_period_days + extension_days_granted.", "Visual warning badges flag overdue deals on both Admin and Affiliate portals."),
        ("Forced Closure by Admin", "Triggered by Admin for overdue or non-responsive attempts.", "Deal status becomes ABORTED with is_force_closed = 1. Company name is unlocked unless appealed."),
        ("Affiliate Appeal for Extension", "Affiliate submits appeal with business justification (e.g. client fiscal budget deferral).", "Status becomes APPEAL_SUBMITTED, pausing name release pending review."),
        ("Admin Appeal Adjudication", "Admin approves (+X days extension) or rejects appeal.", "Approved appeals re-open the deal to PROPOSAL_SENT and grant extra days."),
        ("Direct Admin Extension", "Admin can proactively grant +X days extension at any time without appeal.", "Adds days to extension_days_granted and updates closure deadline.")
    ]

    tbl_closure = doc.add_table(rows=len(closure_table_data), cols=3)
    for r_idx, row in enumerate(closure_table_data):
        for c_idx, text in enumerate(row):
            tbl_closure.cell(r_idx, c_idx).text = text
    style_table(tbl_closure, col_widths=[1.8, 3.2, 2.2])

    # SECTION 5: RATE LOCKING & MANAGEMENT APPROVAL
    add_header_styled(doc, "5. Payment Collection, Snapshot Rate Locking & Management Approval", level=1)
    doc.add_paragraph(
        "When an invoice collection occurs, financial integrity is maintained through a two-step verification and immutability protocol:"
    )

    create_callout(doc, "The Immutability & Rate Locking Protocol",
                   "1. Collection Recording: Admin records invoice #, total amount, collected amount, bank receipt reference, and optional proof media upload (PDF/PNG).\n"
                   "2. Snapshot Capture: Active commission percentages (Direct %, L1 %, L2 %) are captured and written to deal_collections.locked_direct_rate_pct, locked_upline_l1_rate_pct, locked_upline_l2_rate_pct.\n"
                   "3. Gated Status: Created with approval_status = 'PENDING_APPROVAL' and is_immutable = 0.\n"
                   "4. Management Approval: Management verifies cleared bank funds and executes '✓ Acknowledge & Approve Collection'.\n"
                   "5. Immutability Seal: Record is updated to APPROVED, is_immutable = 1, approved_by, and approved_at.\n"
                   "6. Advice Creation: Generates immutable Payment Advice vouchers (ADV-2026-XXXXX) for direct affiliate and uplines.",
                   bg_hex="F0FDF4", border_hex="16A34A")

    # SECTION 6: CONSOLIDATED BATCH DISBURSEMENTS
    add_header_styled(doc, "6. Consolidated Payouts & Batch Disbursement Operations", level=1)
    doc.add_paragraph(
        "To streamline banking operations and eliminate single-transfer overhead, all approved payment advices are consolidated by Beneficiary Affiliate:"
    )

    payout_flow = [
        ("Affiliate Aggregation", "The Admin Payouts dashboard groups all pending vouchers (direct commissions and upline overrides) per affiliate."),
        ("1-Click Single Transfer", "Admin executes one consolidated bank transfer for the aggregate pending balance."),
        ("Single Transaction Reference", "Admin logs the transaction reference (e.g. MBB-9831902) and bank account details."),
        ("Master Batch Voucher (DISB-2026-XXXXX)", "A master record in payout_batches is created, and all constituent payment advices are atomically marked as PAID and linked to payout_batch_id."),
        ("Affiliate Portal Receipt", "Affiliates view their Consolidated Disbursement Receipt alongside full itemized transaction breakdowns.")
    ]
    for title, desc in payout_flow:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(3)
        r_t = bp.add_run(f"{title}: ")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(15, 118, 110)
        r_d = bp.add_run(desc)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # SECTION 7: RELATIONAL DATABASE SCHEMA
    add_header_styled(doc, "7. Relational Database Schema & Data Dictionary", level=1)
    
    db_tables = [
        ("Table Name", "Primary Key", "Key Foreign Keys", "Description & Columns"),
        ("system_settings", "setting_key", "None", "Global rates and closure policies: direct_commission_rate_pct, upline_l1_commission_rate_pct, upline_l2_commission_rate_pct, default_closure_period_days."),
        ("deal_pipeline", "id", "affiliate_id -> affiliate_applications(id)", "Core deal entity: deal_code, customer_name, customer_email, customer_phone, package_name, contract_value_myr, extension_days_granted, is_force_closed, force_closed_at, force_closed_reason, appeal_status, appeal_reason, appeal_submitted_at, appeal_adjudicated_at, appeal_adjudication_notes, status, invoice_number."),
        ("deal_funnel_steps", "id", "deal_id -> deal_pipeline(id)\nsubmitted_by -> users(id)\nreviewed_by -> users(id)", "Milestone progression logs: from_stage, to_stage, step_title, affiliate_notes, submitted_at, admin_review_status (PENDING_REVIEW, ACKNOWLEDGED, RETURNED_FOR_REVIEW), admin_remarks, reviewed_at, is_immutable."),
        ("deal_closure_logs", "id", "deal_id -> deal_pipeline(id)\nperformed_by -> users(id)", "Audit trail of closure actions: action_type (FORCED_CLOSURE, APPEAL_SUBMITTED, EXTENSION_GRANTED, APPEAL_REJECTED), days_extended, action_notes, performed_by_user_id, created_at."),
        ("deal_collections", "id", "deal_id -> deal_pipeline(id)\napproved_by -> users(id)", "Invoice collections: invoice_number, invoice_total_myr, collected_amount_myr, bank_receipt_ref, proof_media_path, locked_direct_rate_pct, locked_upline_l1_rate_pct, locked_upline_l2_rate_pct, approval_status, approved_by, approved_at, approval_remarks, is_immutable."),
        ("payout_batches", "id", "beneficiary_affiliate_id -> affiliate_applications(id)\ndisbursed_by -> users(id)", "Consolidated disbursement batches: batch_code (DISB-2026-XXXXX), total_amount_myr, advice_count, manual_bank_tx_ref, bank_name, bank_account_number, payout_notes, disbursed_by, disbursed_at."),
        ("payment_advices", "id", "collection_id -> deal_collections(id)\ndeal_id -> deal_pipeline(id)\nbeneficiary_affiliate_id -> affiliate_applications(id)\npayout_batch_id -> payout_batches(id)", "Individual commission vouchers: advice_number (ADV-2026-XXXXX), beneficiary_type, rate_percentage, collection_amount_base_myr, commission_amount_myr, payout_status, paid_at, manual_bank_tx_ref, payout_batch_id, is_immutable.")
    ]

    tbl_db = doc.add_table(rows=len(db_tables), cols=4)
    for r_idx, row in enumerate(db_tables):
        for c_idx, text in enumerate(row):
            tbl_db.cell(r_idx, c_idx).text = text
    style_table(tbl_db, col_widths=[1.5, 1.1, 2.2, 2.4])

    # SECTION 8: API ROUTE ARCHITECTURE
    add_header_styled(doc, "8. API Route Endpoints & Integration Map", level=1)

    api_endpoints = [
        ("HTTP Route", "Method", "Auth Guard", "Operational Behavior"),
        ("/api/admin/settings", "POST", "Admin", "Updates commission rates and default closure period days."),
        ("/api/admin/deals/[id]/step-review", "POST", "Admin", "Acknowledges stage update or returns step with mandatory remarks."),
        ("/api/admin/deals/[id]/closure", "POST", "Admin", "Executes force closure, direct extensions (+X days), or adjudicates affiliate appeals."),
        ("/api/admin/deals/[id]/collect", "POST", "Admin", "Submits invoice collection with proof upload and locks snapshot rates."),
        ("/api/admin/collections/[id]/approve", "POST", "Admin", "Executes management approval and generates immutable payment advices."),
        ("/api/admin/payouts/batch", "POST", "Admin", "Executes 1-click consolidated batch disbursement for an affiliate."),
        ("/api/portal/prospects", "POST", "Affiliate", "Names new prospect with anti-poaching exclusivity validation."),
        ("/api/portal/prospects/[id]/step", "POST", "Affiliate", "Submits milestone funnel update or resubmits returned step."),
        ("/api/portal/prospects/[id]/appeal", "POST", "Affiliate", "Submits extension appeal for force-closed prospect.")
    ]

    tbl_api_map = doc.add_table(rows=len(api_endpoints), cols=4)
    for r_idx, row in enumerate(api_endpoints):
        for c_idx, text in enumerate(row):
            tbl_api_map.cell(r_idx, c_idx).text = text
    style_table(tbl_api_map, col_widths=[2.3, 0.8, 1.0, 3.1])

    # SECTION 9: USER INTERFACE ARCHITECTURE
    add_header_styled(doc, "9. User Interface Layout & Screen Specifications", level=1)
    
    doc.add_heading("Affiliate Portal Interfaces (/portal/*)", level=2)
    p_views = [
        ("Prospects Directory (/portal/prospects)", "Comprehensive directory of all introduced prospects with active stage badges, closure deadline countdowns, overdue alerts, and '+ Name & Register New Prospect' modal."),
        ("Prospect Funnel Journey (/portal/prospects/[id])", "Visual stepper path, countdown progress tracker, chronological history of all steps, admin review status badges, reviewer remarks callouts, 'Log Funnel Step' modal, 'Resubmit Step' modal, and 'Submit Extension Appeal' modal."),
        ("Commissions & Payment Advices (/portal)", "Earnings composition, consolidated disbursement receipts (DISB-2026-XXXXX), and formal voucher notes.")
    ]
    for title, desc in p_views:
        bp = doc.add_paragraph(style='List Bullet')
        r_t = bp.add_run(f"{title}: ")
        r_t.font.bold = True
        r_d = bp.add_run(desc)

    doc.add_heading("Admin Dashboard Interfaces (/admin/*)", level=2)
    a_views = [
        ("Deal Pipeline & Management (/admin/deals)", "Visual board with funnel stage filters, quick progression controls, and links to detailed journey paths."),
        ("Deal Journey & Step Review (/admin/deals/[id])", "Full prospect audit trail, '✓ Acknowledge Step Update' and '↩️ Return for Review' buttons, direct closure deadline extension modal, force closure modal, and appeal adjudication interface."),
        ("Collections & Management Approval (/admin/collections)", "Verification queue for invoice receipts, proof media viewer, and management approval execution."),
        ("Consolidated Disbursements (/admin/payouts)", "Grouped pending balances per affiliate with 1-click batch transfer modal, itemized advice viewer, and historical batch archive."),
        ("Commission & Closure Settings (/admin/settings)", "Configuration of Direct %, Upline L1 %, Upline L2 %, and Default Prospect Closure Period (Days).")
    ]
    for title, desc in a_views:
        bp = doc.add_paragraph(style='List Bullet')
        r_t = bp.add_run(f"{title}: ")
        r_t.font.bold = True
        r_d = bp.add_run(desc)

    # DOCUMENT FOOTER
    doc.add_paragraph("\n" + "―" * 55)
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = footer_p.add_run("FolioDesk Enterprise Architecture Specification | Document Ref: FD-ARCH-2026-V3")
    r_foot.font.name = 'Calibri'
    r_foot.font.italic = True
    r_foot.font.size = Pt(9)
    r_foot.font.color.rgb = RGBColor(148, 163, 184)

    # Save to both target directories
    out_dir1 = r"d:\Firedancerx\OneDrive\Work2026\Mujib\Work"
    out_dir2 = r"d:\Websites\FolioDesk\WebsiteBuild\docs"
    os.makedirs(out_dir1, exist_ok=True)
    os.makedirs(out_dir2, exist_ok=True)
    
    file_name = "FolioDesk_Affiliate_Sales_Funnel_and_Payment_Architecture.docx"
    path1 = os.path.join(out_dir1, file_name)
    path2 = os.path.join(out_dir2, file_name)
    
    doc.save(path1)
    doc.save(path2)
    print(f"Saved DOCX successfully to:\n  1. {path1}\n  2. {path2}")

if __name__ == "__main__":
    build_docx()
